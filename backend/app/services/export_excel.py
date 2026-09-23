"""导出费用归集 Excel（四个 Sheet）。

文件名自动包含柜号；若存在未解决关键异常，禁止使用"最终确认"状态/文件名。
"""
from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .allocation_service import compute, _load
from ..core.decimal_utils import q2
from ..core.volume import ref_volume_ratio

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=12, color="1F4E78")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY_FMT = "0.00"
VOL_FMT = "0.000000"


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _autosize(ws, max_w=40):
    for col in ws.columns:
        length = 0
        letter = None
        for cell in col:
            if cell.value is not None:
                letter = cell.column_letter
                length = max(length, len(str(cell.value)))
        if letter:
            ws.column_dimensions[letter].width = min(max(length + 2, 10), max_w)


def export_session_excel(session_id: str, mode: str = "final") -> bytes:
    s = _load(session_id)
    data = s["data"]
    res = compute(session_id)
    recon = res["reconciliation"]
    refs = res["refs"]
    matches = res["matches"]
    sea_alloc = res["sea_freight_alloc"]
    article_rmb = res["article_rmb"]
    duty_alloc = res["duty_alloc"]
    pending_duty = Decimal(res["pending_duty"])

    cabinet_no = data.get("cabinet_no", "")
    unresolved = recon.get("unresolved_exceptions", False)
    if mode == "final" and unresolved:
        raise ValueError("存在未解决的关键归属异常，禁止导出最终确认版；请先处理待确认税项或导出待确认工作版。")

    # 索引：ref -> 原始 HS / 商品类别
    ref_hs = {}
    ref_category = {}
    for r in refs:
        hss = []
        cats = []
        for it in r.get("items", []):
            if it.get("hs_code"):
                hss.append(it["hs_code"])
            nm = it.get("en_name") or it.get("cn_name") or ""
            if nm:
                cats.append(nm)
        ref_hs[r["ref_id"]] = "/".join(sorted(set(hss)))
        ref_category[r["ref_id"]] = "、".join(sorted(set(cats)))

    # 税项原始数据（计税金额等）
    article_by_key = {}
    for a in data["customs"]["articles"]:
        key = f"{a.get('mrn','')}|{a.get('article_no','')}"
        article_by_key[key] = a

    def _ref_volume(r: dict) -> Decimal:
        vol_raw = r.get("volume_m3")
        if vol_raw in (None, ""):
            vol_raw = sum((Decimal(str(it.get("box_type_volume_m3") or 0))
                           for it in r.get("items", [])), Decimal("0"))
        return Decimal(str(vol_raw))

    ref_vols = {r["ref_id"]: _ref_volume(r) for r in refs}
    total_volume = sum(ref_vols.values()) or Decimal("1")

    wb = openpyxl.Workbook()

    # ---- Sheet1 货件费用归集 ----
    ws1 = wb.active
    ws1.title = "货件费用归集"
    ws1.append(["Reference ID", "商品类别", "箱数", "总体积（m³）", "体积占比",
                "海运费（元）", "关税（元）", "费用合计（元）", "归集状态"])
    _style_header(ws1, 1, 9)
    cum_box = cum_vol = cum_sea = cum_duty = cum_total = 0
    for r in refs:
        rid = r["ref_id"]
        vol = ref_vols[rid]
        box_count = r.get("total_box_count")
        if box_count is None:
            box_count = sum(int(it.get("box_count", 0) or 0) for it in r.get("items", []))
        box_count = int(box_count)
        sea = Decimal(sea_alloc.get(rid, "0"))
        duty = Decimal(duty_alloc.get(rid, "0"))
        ratio = ref_volume_ratio(vol, total_volume)
        # 状态：该货件是否有待确认税项
        ref_has_pending = any(
            m["ref_id"] == rid and m["status"] in ("pending", "unmatched") for m in matches)
        status = "含待确认" if ref_has_pending else "已归集"
        ws1.append([rid, ref_category.get(rid, ""), box_count,
                    float(vol), float(ratio), float(sea), float(duty),
                    float(sea + duty), status])
        cum_box += box_count
        cum_vol += vol
        cum_sea += sea
        cum_duty += duty
        cum_total += sea + duty
    # 合计行
    ws1.append(["合计", "", cum_box, float(cum_vol), 1.0,
                float(cum_sea), float(cum_duty), float(cum_total), ""])
    # 待确认关税（不虚构 Reference ID）
    if pending_duty > 0:
        ws1.append(["待确认-未归属税项", "见关税匹配明细 Sheet", "", "", "",
                    "", float(pending_duty), float(pending_duty), "待确认"])
    for row in ws1.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            if cell.column in (4, 5):
                cell.number_format = VOL_FMT if cell.column == 4 else "0.0000"
            elif cell.column in (6, 7, 8):
                cell.number_format = MONEY_FMT

    # ---- Sheet2 箱型体积明细 ----
    ws2 = wb.create_sheet("箱型体积明细")
    ws2.append(["Reference ID", "SKU", "货箱编号", "箱数", "长度（cm）", "宽度（cm）",
                "高度（cm）", "单箱体积（m³）", "该箱型总体积（m³）"])
    _style_header(ws2, 1, 9)
    for r in refs:
        rid = r["ref_id"]
        for it in r.get("items", []):
            sv = it.get("single_volume_m3") or 0
            bv = it.get("box_type_volume_m3") or 0
            ws2.append([rid, it.get("sku", ""), it.get("box_spec", ""),
                        it.get("box_count", 0), float(it.get("length_cm") or 0),
                        float(it.get("width_cm") or 0), float(it.get("height_cm") or 0),
                        float(sv), float(bv)])
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            if cell.column in (8, 9):
                cell.number_format = VOL_FMT

    # ---- Sheet3 关税匹配明细 ----
    ws3 = wb.create_sheet("关税匹配明细")
    ws3.append(["Reference ID", "SKU/商品名称", "原始HS Code", "报关HS Code",
                "申报数量", "计税金额", "原币关税", "人民币关税", "分摊依据", "匹配状态"])
    _style_header(ws3, 1, 10)
    for m in matches:
        if m.get("ref_id"):
            rid = m["ref_id"]
        elif m.get("split_weights"):
            rid = "合并拆分：" + "/".join(m.get("candidate_refs", []))
        elif m["status"] in ("pending", "unmatched"):
            rid = "待确认：" + "/".join(m.get("candidate_refs", []))
        else:
            rid = "—"
        art = article_by_key.get(m["article_key"], {})
        base = art.get("duty_base")
        rmb = article_rmb.get(m["article_key"], "0")
        ws3.append([
            rid, m.get("description", ""),
            ref_hs.get(m.get("ref_id"), "") if m.get("ref_id") else "",
            m.get("hs_code", ""),
            float(m["declared_qty"]) if m.get("declared_qty") is not None else "",
            float(base) if base not in (None, "") else "",
            float(m.get("duty_eur", 0)), float(rmb),
            m.get("reason", ""), m.get("status", ""),
        ])
    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            if cell.column in (6, 7, 8):
                cell.number_format = MONEY_FMT

    # ---- Sheet4 核对与异常 ----
    ws4 = wb.create_sheet("核对与异常")
    rows4 = [
        ("原始海运费（元）", recon["original_sea_freight"]),
        ("分摊后海运费（元）", recon["allocated_sea_freight"]),
        ("海运费差额（元）", recon["sea_freight_diff"]),
        ("原币关税合计（€）", recon["eur_duty_total"]),
        ("结算汇率", recon["exchange_rate"]),
        ("实际人民币关税（元）", recon["rmb_duty_input"]),
        ("已归集关税（元）", recon["allocated_duty"]),
        ("待确认关税（元）", recon["pending_duty"]),
        ("关税核对差额（元）", recon["duty_diff"]),
        ("装箱单总箱数", recon["box_count_packing"]),
        ("海关申报总件数（colli）", recon["box_count_customs"] if recon["box_count_customs"] is not None else "—"),
        ("箱数差异", recon["box_count_diff"] if recon["box_count_diff"] is not None else "—"),
    ]
    ws4.append(["核对项目", "数值"])
    _style_header(ws4, 1, 2)
    for k, v in rows4:
        ws4.append([k, float(v) if isinstance(v, (int, float, Decimal)) else v])
    for row in ws4.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            if isinstance(cell.value, float):
                cell.number_format = MONEY_FMT

    ws4.append([])
    ws4.append(["报关数量差异", "", "", ""])
    ws4.append(["Reference ID", "税项", "海关申报数量", "装箱数量", "差异"])
    _style_header(ws4, ws4.max_row, 5)
    for d in recon.get("declared_qty_diff", []):
        ws4.append([d["ref_id"], d["article"], float(d["customs_qty"]),
                    float(d["packing_qty"]), float(d["diff"])])
    ws4.append([])
    ws4.append(["HS Code 异常", "", "", ""])
    ws4.append(["Reference ID", "税项", "报关HS", "装箱HS", "说明"])
    _style_header(ws4, ws4.max_row, 5)
    for h in recon.get("hs_anomalies", []):
        ws4.append([h["ref_id"], h["article"], h["customs_hs"],
                    "/".join(h["packing_hs"]), h["note"]])
    ws4.append([])
    ws4.append(["尾差处理说明", recon.get("tail_handling", "")])
    ws4.append(["未解决关键异常", "是" if recon.get("unresolved_exceptions") else "否"])

    for ws in (ws1, ws2, ws3, ws4):
        _autosize(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def build_export_filename(cabinet_no: str, mode: str) -> str:
    suffix = "最终分摊明细" if mode == "final" else "待确认分摊明细"
    return f"{cabinet_no}大柜_海运费及关税费用归集_{suffix}.xlsx"
