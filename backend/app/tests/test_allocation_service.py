from decimal import Decimal
import io

import openpyxl

from app.services import allocation_service as svc
from app.services.export_excel import export_session_excel
from app.tests.conftest import (build_packing_excel_008, build_uitnodiging_pdf_008,
                                build_packing_excel_009, build_uitnodiging_pdf_009)


def test_008_end_to_end(tmp_db):
    sid = svc.create_session("008", 54485.80, 7486.68, 8.2, "回归")
    svc.parse_and_store_packing(sid, build_packing_excel_008())
    svc.parse_and_store_customs(sid, [("njeu.pdf", build_uitnodiging_pdf_008(), True)])
    res = svc.compute(sid)
    sea_sum = sum(Decimal(v) for v in res["sea_freight_alloc"].values())
    assert sea_sum == Decimal("54485.80")
    duty_sum = Decimal(res["reconciliation"]["allocated_duty"]) + Decimal(res["pending_duty"])
    assert duty_sum == Decimal("7486.68")
    # 新规则：HS 不一致不再判待确认；008 全部按名称+数量自动匹配，无待确认
    assert Decimal(res["pending_duty"]) == 0
    assert res["reconciliation"]["unresolved_exceptions"] is False
    assert Decimal(res["reconciliation"]["allocated_duty"]) == Decimal("7486.68")


def test_export_sheet1_has_box_volume_ratio(tmp_db):
    """回归：导出 Sheet1 的箱数/总体积/体积占比不得为 0（此前 property 未序列化导致全 0）。"""
    sid = svc.create_session("008", 54485.80, 7486.68, 8.2, "回归")
    svc.parse_and_store_packing(sid, build_packing_excel_008())
    svc.parse_and_store_customs(sid, [("njeu.pdf", build_uitnodiging_pdf_008(), True)])
    svc.compute(sid)
    wb = openpyxl.load_workbook(io.BytesIO(export_session_excel(sid, "final")))
    ws = wb["货件费用归集"]
    header = [c.value for c in ws[1]]
    i_box = header.index("箱数")
    i_vol = header.index("总体积（m³）")
    i_ratio = header.index("体积占比")
    i_sea = header.index("海运费（元）")
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    data_rows = [r for r in rows if r[0] not in ("合计", "待确认-未归属税项", None)]
    assert data_rows, "应有货件数据行"
    # 箱数合计 610、总体积 > 0、各行占比 > 0 且合计≈1
    assert sum(r[i_box] for r in data_rows) == 610
    assert sum(float(r[i_vol]) for r in data_rows) > 0
    assert all(float(r[i_ratio]) > 0 for r in data_rows)
    assert abs(sum(float(r[i_ratio]) for r in data_rows) - 1.0) < 1e-6
    # 海运费合计仍等于输入
    assert abs(sum(float(r[i_sea]) for r in data_rows) - 54485.80) < 0.005


def test_009_mirror_even_split(tmp_db):
    """009 镜子：两个货件数量相同无法区分 -> 自动按数量均摊，不再挂待确认。"""
    sid = svc.create_session("009", 56965.80, 12284.50, 8.2, "回归")
    svc.parse_and_store_packing(sid, build_packing_excel_009())
    svc.parse_and_store_customs(sid, [("009.pdf", build_uitnodiging_pdf_009(), True)])
    res = svc.compute(sid)
    sea_sum = sum(Decimal(v) for v in res["sea_freight_alloc"].values())
    assert sea_sum == Decimal("56965.80")
    duty_sum = Decimal(res["reconciliation"]["allocated_duty"]) + Decimal(res["pending_duty"])
    assert duty_sum == Decimal("12284.50")
    # 镜子税项自动均摊到两个货件，无待确认
    mirror = next(m for m in res["matches"] if m["description"] == "mirror")
    assert mirror["status"] == "auto"
    assert set(mirror["split_weights"].keys()) == {
        "RVG15056-260909-0001", "RVG15056-260909-0002"}
    assert Decimal(res["pending_duty"]) == 0
    # 两个货件各分到一半镜子关税（税项 1 duty 400€ / 税项 2 lamp：见 conftest）
    d1 = Decimal(res["duty_alloc"]["RVG15056-260909-0001"])
    d2 = Decimal(res["duty_alloc"]["RVG15056-260909-0002"])
    assert d1 == d2 and d1 > 0
